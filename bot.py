import os
import fitz
import logging
import sqlite3
import pytesseract
from PIL import Image
import re
import rarfile
import zipfile
import speech_recognition as sr
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Updater, CommandHandler, MessageHandler, Filters,
    CallbackContext, CallbackQueryHandler
)
from docx import Document
from datetime import datetime
import openpyxl
import requests
from bs4 import BeautifulSoup
from pydub import AudioSegment
from speech_recognition import Recognizer, AudioFile
from googletrans import Translator

def is_suspicious_file(path):
    suspicious_keywords = ["keylogger", "reverse", "vnc", "socket", "payload", "backdoor", "rat", ".exe", ".bat", ".ps1"]
    try:
        with open(path, "rb") as f:
            content = f.read(5000).decode(errors='ignore').lower()
            for keyword in suspicious_keywords:
                if keyword in content:
                    return True
    except:
        return True
    return False

def is_suspicious_url(url):
    blocked = ["phish", "porn", "trojan", "malware", "hacker", "scam", "keylogger", "xn--"]
    return any(bad in url.lower() for bad in blocked)

def clean_filename(name):
    return re.sub(r'[<>:"/\\|?*\x00-\x1F]', '_', name)

rate_limit_map = {}
def is_rate_limited(user_id, delay=2):
    from time import time
    now = time()
    if user_id in rate_limit_map and now - rate_limit_map[user_id] < delay:
        return True
    rate_limit_map[user_id] = now
    return False

BOT_TOKEN = "8085895393:AAEesSGmVM5JgS8-U7odLnm55Z8yoN9gWLo"
pytesseract.pytesseract.tesseract_cmd = r"C:\Users\user\Desktop\AhmadAIbot\tesseract.exe"

logging.basicConfig(format='%(asctime)s - %(name)s - %(levelname)s - %(message)s', level=logging.INFO)
last_extracted_text = {}
user_question_queues = {}  # لتخزين قائمة الأسئلة لكل مستخدم
global_rate_limit = {}


def is_global_rate_limited(delay=1):
    from time import time
    now = time()
    key = int(now)  # لكل ثانية
    if global_rate_limit.get(key, 0) > 20:
        return True
    global_rate_limit[key] = global_rate_limit.get(key, 0) + 1
    return False

def help_command(update: Update, context: CallbackContext):
    text = (
        "🤖 *مرحباً بك في البوت الذكي AhmadAIbot!*\n\n"
        "💡 إليك قائمة بالأوامر والميزات المتاحة:\n\n"
        "📂 /start - مقدمة عن البوت\n"
        "🌐 /readlink [رابط] - قراءة محتوى صفحة ويب وتلخيصه\n"
        "🌍 /translate [اللغة] [النص] - ترجمة نص (مثل: /translate ar Hello)\n"
        "🧠 إرسال ملف (PDF, Word, Excel, صورة, صوت) لتحليله وتلخيصه\n"
        "🔍 إرسال رابط تلقائيًا يقوم بالكشف عنه ويعطيك خيارات (معاينة، تلخيص، ترجمة، فحص أمان)\n\n"
        "👇 اختر ما ترغب بفعله:"
    )
    keyboard = [
        [InlineKeyboardButton("📄 رفع ملف", switch_inline_query_current_chat="")],
        [InlineKeyboardButton("🌐 قراءة رابط", callback_data="show_readlink")],
        [InlineKeyboardButton("🌍 ترجمة", callback_data="show_translate")],
        [InlineKeyboardButton("🧠 مساعدة إضافية", url="https://t.me/your_support_group")]
    ]
    update.message.reply_text(text, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(keyboard))



def init_db():
    conn = sqlite3.connect("lecture_history.db")
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        username TEXT,
        filename TEXT,
        fulltext TEXT,
        summary TEXT,
        questions TEXT,
        timestamp TEXT
    )''')
    conn.commit()
    conn.close()



def extract_rar_file(rar_path):
    extracted = []
    try:
        with rarfile.RarFile(rar_path) as rf:
            rf.extractall("unzipped_rar")
            for f in rf.namelist():
                full_path = os.path.join("unzipped_rar", f)
                if is_suspicious_file(full_path):
                    os.remove(full_path)
                else:
                    extracted.append(full_path)
    except Exception as e:
        print("خطأ بفك RAR:", e)
    return extracted


def is_suspicious_file(path):
    suspicious_keywords = ["keylogger", "reverse", "vnc", "socket", "payload", "backdoor", "rat", ".exe", ".bat", ".ps1"]
    try:
        with open(path, "rb") as f:
            content = f.read(5000).decode(errors='ignore').lower()
            for keyword in suspicious_keywords:
                if keyword in content:
                    return True
    except:
        return True  # إذا ما قدر يقرأه، اعتبره مشبوه
    return False



def extract_and_scan_zip(zip_path):
    extracted = []
    try:
        with zipfile.ZipFile(zip_path, 'r') as zip_ref:
            zip_ref.extractall("unzipped")
            for file in zip_ref.namelist():
                full_path = os.path.join("unzipped", file)
                if is_suspicious_file(full_path):
                    os.remove(full_path)
                else:
                    extracted.append(full_path)
    except Exception as e:
        print("خطأ بفك الضغط:", e)
    return extracted


def is_url(text):
    return re.match(r'^https?://', text.strip())

def get_url_text(url):
    try:
        r = requests.get(url, timeout=10)
        soup = BeautifulSoup(r.text, "html.parser")
        return soup.get_text(separator="\n").strip()
    except:
        return "❌ فشل في جلب المحتوى."

def handle_text(update: Update, context: CallbackContext):
    if is_global_rate_limited():
     return  # تجاهل الرسالة بصمت لو كان فيه ضغط عالمي

    user_id = update.effective_user.id

    # ✅ فحص المعدل
    if is_rate_limited(user_id):
        update.message.reply_text("⏳ الرجاء الانتظار قليلاً قبل المحاولة مجددًا.")
        return

    if user_id in user_question_queues and user_question_queues[user_id]:
        handle_answer(update, context)
        return

    text = update.message.text.strip()
    if is_url(text):
        if is_suspicious_url(text):
            update.message.reply_text("⚠️ تم حظر هذا الرابط لأنه قد يكون ضارًا.")
            return

        context.user_data["last_url"] = text
        update.message.reply_text(
            "📎 تم اكتشاف رابط، اختر ما تريد فعله:",
            reply_markup=InlineKeyboardMarkup([
                [InlineKeyboardButton("🔍 معاينة", callback_data="preview_url")],
                [InlineKeyboardButton("🧠 تلخيص", callback_data="summarize_url")],
                [InlineKeyboardButton("🌍 ترجمة", callback_data="translate_url")],
                [InlineKeyboardButton("⚠️ فحص الرابط", callback_data="check_url")]
            ])
        )

def start(update: Update, context: CallbackContext):
    text = (
        "👋 مرحباً بك في *AhmadAIbot* الذكي!\n\n"
        "💡 هذا البوت يساعدك على:\n"
        "- قراءة وتحليل الملفات (PDF, Word, Excel, صور، صوت)\n"
        "- تلخيص النصوص وطرح الأسئلة\n"
        "- قراءة وتحليل الروابط\n"
        "- الترجمة بين العربية والإنجليزية\n\n"
        "👇 اختر ما ترغب بفعله:"
    )
    keyboard = [
        [InlineKeyboardButton("📂 رفع ملف", switch_inline_query_current_chat="")],
        [InlineKeyboardButton("🌐 قراءة رابط", callback_data="show_readlink")],
        [InlineKeyboardButton("🌍 ترجمة", callback_data="show_translate")],
        [InlineKeyboardButton("🧠 تعليمات / تعليمية", callback_data="show_help")]
    ]
    update.message.reply_text(text, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(keyboard))

def send_back_to_main_menu(message):
    keyboard = [[InlineKeyboardButton("🔙 رجوع إلى القائمة الرئيسية", callback_data="back_to_main")]]
    message.reply_text("⬅️ اختر من القائمة:", reply_markup=InlineKeyboardMarkup(keyboard))


def start_quiz_session(callback_query, context: CallbackContext):
    user_id = callback_query.from_user.id
    if user_id not in user_question_queues or not user_question_queues[user_id]:
        callback_query.message.reply_text("❗ لا توجد أسئلة جاهزة للمراجعة.")
        return

    question, _ = user_question_queues[user_id][0]
    callback_query.message.reply_text(f"❓ *السؤال الأول:*\n\n{question}", parse_mode="Markdown")



def handle_document(update: Update, context: CallbackContext):
    file = update.message.document
    file_name = clean_filename(file.file_name)  # ✅ تنظيف الاسم
    file_id = file.file_id
    new_file = context.bot.get_file(file_id)
    file_path = f"downloads/{file_name}"
    os.makedirs("downloads", exist_ok=True)
    new_file.download(file_path)

    if is_suspicious_file(file_path):
        update.message.reply_text("⚠️ تم حظر هذا الملف لأنه يحتوي على محتوى مشبوه.")
        os.remove(file_path)
        return

    try:
        ext = file_name.lower()
        extracted_paths = None

        # ✅ فك ضغط الملفات المضغوطة
        if ext.endswith(".zip"):
            extracted_paths = extract_and_scan_zip(file_path)
        elif ext.endswith(".rar"):
            extracted_paths = extract_rar_file(file_path)

        # ✅ معالجة الملفات المستخرجة
        if extracted_paths is not None:
            if not extracted_paths:
                update.message.reply_text("❌ لم يتم استخراج ملفات أو تم حذف ملفات ضارة.")
                return

            msg = "📦 تم استخراج الملفات التالية:\n"
            for path in extracted_paths:
                msg += f"- {os.path.basename(path)}\n"
                try:
                    with open(path, "rb") as f:
                        update.message.reply_document(f, filename=os.path.basename(path))
                except:
                    update.message.reply_text(f"❌ فشل في إرسال: {os.path.basename(path)}")

            update.message.reply_text(msg)

            for path in extracted_paths:
                text = ""
                if path.endswith(".pdf"):
                    text = extract_text_from_pdf(path)
                elif path.endswith(".docx"):
                    text = extract_text_from_word(path)
                elif path.endswith(".xlsx"):
                    text = extract_text_from_excel(path)
                elif path.endswith((".png", ".jpg", ".jpeg")):
                    text = extract_text_from_image(path)
                elif path.endswith((".wav", ".mp3", ".opus")):
                    text = extract_text_from_audio(path)

                if not text.strip():
                    continue

                update.message.reply_text(f"📤 تم استخراج النص من `{os.path.basename(path)}`:", parse_mode="Markdown")
                for i in range(0, len(text), 4000):
                    update.message.reply_text(text[i:i+4000])

                summary = simple_summary(text)
                questions = extract_questions(text)

                user_question_queues[update.effective_user.id] = questions

                if questions:
                    update.message.reply_text(
                        "🧠 تم توليد أسئلة مراجعة من النص. هل ترغب ببدء جلسة مراجعة؟",
                        reply_markup=InlineKeyboardMarkup([
                            [InlineKeyboardButton("ابدأ المراجعة ✅", callback_data="start_quiz")]
                        ])
                    )

                save_to_db(update.effective_user.id, update.effective_user.username, os.path.basename(path), text, summary, questions)

                reply = f"*📄 الملخص:*\n\n{summary}\n\n*❓ الأسئلة:*\n" + "\n".join([f"- {q}" for q in questions])
                for i in range(0, len(reply), 4000):
                    update.message.reply_text(reply[i:i+4000], parse_mode="Markdown")

                send_translate_buttons(update)

            send_back_to_main_menu(update.message)
            return

        # ✅ ملفات غير مضغوطة
        text = ""
        if ext.endswith(".pdf"):
            text = extract_text_from_pdf(file_path)
        elif ext.endswith(".docx"):
            text = extract_text_from_word(file_path)
        elif ext.endswith(".xlsx"):
            text = extract_text_from_excel(file_path)
        elif ext.endswith((".png", ".jpg", ".jpeg")):
            text = extract_text_from_image(file_path)
        elif ext.endswith((".wav", ".mp3", ".opus")):
            text = extract_text_from_audio(file_path)
        else:
            update.message.reply_text("❌ نوع الملف غير مدعوم.")
            return

        if not text.strip():
            update.message.reply_text("❌ لم يتم استخراج نص.")
            return

        last_extracted_text[update.effective_user.id] = text
        update.message.reply_text("📤 تم استخراج النص الكامل:")
        for i in range(0, len(text), 4000):
            update.message.reply_text(text[i:i+4000])

        summary = simple_summary(text)
        questions = extract_questions(text)
        user_question_queues[update.effective_user.id] = questions

        if questions:
            update.message.reply_text(
                "🧠 تم توليد أسئلة مراجعة من النص. هل ترغب ببدء جلسة مراجعة؟",
                reply_markup=InlineKeyboardMarkup([
                    [InlineKeyboardButton("ابدأ المراجعة ✅", callback_data="start_quiz")]
                ])
            )

        save_to_db(update.effective_user.id, update.effective_user.username, file_name, text, summary, questions)

        reply = f"*📄 الملخص المستخرج:*\n\n{summary}\n\n*❓ الأسئلة:*\n" + "\n".join([f"- {q}" for q in questions])
        for i in range(0, len(reply), 4000):
            update.message.reply_text(reply[i:i+4000], parse_mode="Markdown")

        send_translate_buttons(update)

        # ✅ زر ضغط الملف
        if not ext.endswith((".zip", ".rar")):
            keyboard = [[InlineKeyboardButton("📦 ضغط الملف", callback_data=f"compress:{file_path}")]]
            update.message.reply_text("هل ترغب بضغط الملف؟", reply_markup=InlineKeyboardMarkup(keyboard))

    except Exception as e:
        update.message.reply_text(f"❌ خطأ: {str(e)}")







         

def handle_answer(update: Update, context: CallbackContext):
    user_id = update.effective_user.id
    if user_id not in user_question_queues or not user_question_queues[user_id]:
        return  # لا يوجد جلسة نشطة

    user_question_queues[user_id].pop(0)

    if not user_question_queues[user_id]:
        update.message.reply_text("✅ انتهت جلسة المراجعة. أحسنت 👏")
        return

    next_q, _ = user_question_queues[user_id][0]
    update.message.reply_text(f"❓ *السؤال التالي:*\n\n{next_q}", parse_mode="Markdown")



def handle_photo(update: Update, context: CallbackContext):
    photo = update.message.photo[-1].get_file()
    file_path = f"downloads/photo_{update.message.message_id}.jpg"
    os.makedirs("downloads", exist_ok=True)
    photo.download(file_path)

    try:
        text = extract_text_from_image(file_path)
        if not text.strip():
            update.message.reply_text("❌ لا يوجد نص واضح.")
            return

        last_extracted_text[update.effective_user.id] = text
        update.message.reply_text("📤 النص المستخرج من الصورة:")
        update.message.reply_text(text)

        summary = simple_summary(text)
        questions = extract_questions(text)
        reply = f"*📄 الملخص:*\n\n{summary}\n\n*❓ الأسئلة:*\n" + "\n".join([f"- {q}" for q in questions])
        update.message.reply_text(reply, parse_mode="Markdown")

        save_to_db(update.effective_user.id, update.effective_user.username, "photo.jpg", text, summary, questions)
        send_translate_buttons(update)


    except Exception as e:
        update.message.reply_text(f"❌ خطأ أثناء المعالجة: {str(e)}")

def extract_text_from_pdf(path):
    doc = fitz.open(path)
    return "".join([page.get_text() for page in doc])

def extract_text_from_word(path):
    doc = Document(path)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_from_excel(path):
    wb = openpyxl.load_workbook(path)
    rows = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            rows.append(" | ".join(str(cell) if cell else "" for cell in row))
    return "\n".join(rows)

def extract_text_from_image(path):
    with open(path, 'rb') as f:
        response = requests.post(
            'https://api.ocr.space/parse/image',
            files={'filename': f},
            data={'apikey': 'YOUR_OCR_SPACE_API_KEY'}
        )
    return response.json()['ParsedResults'][0]['ParsedText']


def extract_text_from_audio(path):
    recognizer = Recognizer()
    wav_path = path.rsplit(".", 1)[0] + ".wav"
    AudioSegment.from_file(path).export(wav_path, format="wav")
    with AudioFile(wav_path) as source:
        audio = recognizer.record(source)
        return recognizer.recognize_google(audio, language="ar-JO")

def simple_summary(text):
    lines = text.splitlines()
    return "\n".join([l for l in lines if len(l.strip()) > 30][:5]) or "لا يوجد ملخص واضح."

def extract_questions(text):
    return [(l.strip(), "") for l in text.splitlines() if "?" in l][:5]


def save_to_db(uid, username, fname, fulltext, summary, questions):
    conn = sqlite3.connect("lecture_history.db")
    c = conn.cursor()
    c.execute("INSERT INTO history (user_id, username, filename, fulltext, summary, questions, timestamp) VALUES (?, ?, ?, ?, ?, ?, ?)",
              (uid, username, fname, fulltext, summary, "\n".join(questions), datetime.now().isoformat()))
    conn.commit()
    conn.close()

def readlink(update: Update, context: CallbackContext):
    if not context.args:
        update.message.reply_text("❗ أرسل الرابط بعد الأمر: /readlink https://...")
        return
    url = context.args[0]
    try:
        html = requests.get(url).text
        soup = BeautifulSoup(html, "html.parser")
        text = soup.get_text(separator="\n")
        last_extracted_text[update.effective_user.id] = text
        update.message.reply_text("📄 تم استخراج النص من الرابط.")
        update.message.reply_text(simple_summary(text))
        send_translate_buttons(update)
    except Exception as e:
        update.message.reply_text(f"❌ خطأ: {e}")


def send_translate_buttons(update: Update):
    keyboard = [
        [InlineKeyboardButton("🇸🇦 ترجمة للعربية", callback_data="translate_ar")],
        [InlineKeyboardButton("🇬🇧 Translate to English", callback_data="translate_en")]
    ]
    reply_markup = InlineKeyboardMarkup(keyboard)
    update.message.reply_text("🌐 هل ترغب بترجمة النص؟", reply_markup=reply_markup)

def handle_translate_command(update: Update, context: CallbackContext):
    if not context.args:
        update.message.reply_text("❗ استخدم الأمر: /translate ar Hello world")
        return
    lang = context.args[0]
    text = " ".join(context.args[1:])
    if not text:
        text = last_extracted_text.get(update.effective_user.id, "")
        if not text:
            update.message.reply_text("❌ لا يوجد نص محفوظ لترجمته.")
            return
    translator = Translator()
    translated = translator.translate(text, dest=lang)
    update.message.reply_text(f"🌍 الترجمة:\n{translated.text}")


    


def handle_callback(update: Update, context: CallbackContext):
    query = update.callback_query
    user_id = query.from_user.id
    data = query.data

    if data in ["translate_ar", "translate_en"]:
        lang = "ar" if data == "translate_ar" else "en"
        text = last_extracted_text.get(user_id, "")
        if not text:
            query.answer("❌ لا يوجد نص.")
            return
        translator = Translator()
        translated = translator.translate(text, dest=lang).text
        context.user_data["last_translated_text"] = translated
        keyboard = [[InlineKeyboardButton("🧠 تلخيص الترجمة", callback_data="summarize_translated")]]
        query.message.reply_text(f"🌍 الترجمة:\n{translated}", reply_markup=InlineKeyboardMarkup(keyboard))
        query.answer()

    elif data == "summarize_translated":
        translated_text = context.user_data.get("last_translated_text", "")
        if not translated_text:
            query.answer("❌ لا يوجد ترجمة محفوظة.")
            return
        summary = simple_summary(translated_text)
        query.message.reply_text(f"🧠 *ملخص الترجمة:*\n\n{summary}", parse_mode="Markdown")
        query.answer()

    elif data == "preview_url":
        url = context.user_data.get("last_url", "")
        content = get_url_text(url)
        query.message.reply_text(f"🔍 المحتوى:\n\n{content[:4000]}")

    elif data == "summarize_url":
        url = context.user_data.get("last_url", "")
        content = get_url_text(url)
        summary = simple_summary(content)
        query.message.reply_text(f"🧠 *الملخص:*\n\n{summary}", parse_mode="Markdown")

    elif data == "translate_url":
        url = context.user_data.get("last_url", "")
        content = get_url_text(url)
        translated = Translator().translate(content, dest="ar").text
        query.message.reply_text(f"🌍 *الترجمة:*\n\n{translated[:4000]}", parse_mode="Markdown")

    elif data == "check_url":
        url = context.user_data.get("last_url", "")
        if any(bad in url.lower() for bad in ["hacker", "phish", "porn", "malware", "virus", "xn--"]):
            query.message.reply_text("⚠️ الرابط يحتوي على كلمات مشبوهة وقد يكون ضارًا.")
        else:
            query.message.reply_text("✅ الرابط يبدو آمنًا بناءً على فحص سريع.")
        query.answer()


    elif data == "start_quiz":
        start_quiz_session(update.callback_query, context)
        update.callback_query.answer()


    elif data.startswith("compress:"):
     path = data.split(":", 1)[1]
     zip_path = path + ".zip"
     try:
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            zipf.write(path, os.path.basename(path))
        with open(zip_path, "rb") as f:
            query.message.reply_document(f, filename=os.path.basename(zip_path))
     except Exception as e:
        query.message.reply_text(f"❌ خطأ أثناء ضغط الملف: {str(e)}")
        query.answer()
    

    elif data == "back_to_main":
     query.answer("⬅️ عدت إلى القائمة الرئيسية")


def main():
    init_db()
    updater = Updater(BOT_TOKEN, use_context=True)
    dp = updater.dispatcher
    dp.add_handler(CommandHandler("start", start))
    dp.add_handler(CommandHandler("readlink", readlink))
    dp.add_handler(CommandHandler("translate", handle_translate_command))
    dp.add_handler(CallbackQueryHandler(handle_callback))
    dp.add_handler(MessageHandler(Filters.document, handle_document))
    dp.add_handler(MessageHandler(Filters.photo, handle_photo))
    dp.add_handler(MessageHandler(Filters.text & ~Filters.command, handle_text))
    dp.add_handler(CommandHandler("help", help_command))
    updater.start_polling()
    updater.idle()
    


if __name__ == '__main__':
    main()

